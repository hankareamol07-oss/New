import sqlite3,json,random,sys,html,asyncio,re,os,base64
from playwright.async_api import async_playwright
from docx import Document
from docx.shared import Pt,Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from gen_paper import DB,FONT,LOGO,MR,STD,DEV,INS,ALT,goodpairs
FIG='/home/ubuntu/repos/New/exam_paper/data/book_figures/'
LETTERS='अ ब क ड इ ई उ ऊ'.split()
# plan per total: list of (qtype, n, marks each)
PLANS={20:[('fill_blank',4,1),('mcq',3,1),('true_false',3,1),('figure',3,1),('one_word',3,1),('one_sentence',2,2)],
 30:[('fill_blank',5,1),('true_false',4,1),('mcq',4,1),('figure',3,1),('one_word',4,1),('match',1,4),('one_sentence',3,2)],
 40:[('fill_blank',5,1),('true_false',5,1),('one_sentence',5,1),('mcq',5,1),('figure',3,1),('match',1,4),('one_word',3,1),('short_answer',3,2),('reason',1,3)],
 50:[('fill_blank',5,1),('true_false',5,1),('one_sentence',5,1),('mcq',5,1),('figure',4,1),('match',1,5),('one_word',4,1),('short_answer',4,2),('reason',3,3)]}
BAD="and text not like '%चित्र%' and text not like '%picture%' and text not like '%देखकर%' and text not like '%पढ़कर%' and text not like '%वाचून%' and text not like '%उतारा%' and text not like '%बघून%'"
def pick(db,std,sub,med,total,seed=27):
    random.seed(f'{std}{sub}{med}{seed}')
    chs=db.execute("select c.chapter_id,c.chapter_no,c.title from ep_book_chapters c join ep_books b on b.book_id=c.book_id where b.standard=? and b.subject=? and b.medium=? order by c.chapter_no",(std,sub,med)).fetchall()
    chs=chs[:max(3,round(len(chs)*0.5))]; ids=[c[0] for c in chs]; ph=','.join('?'*len(ids))
    q=f"select bq_id,qtype,text,options_json,answer,pairs_json,figure_image from ep_book_questions where chapter_id in ({ph}) and qtype=? and length(text)>6 and length(text)<220 {BAD}"
    used=set();secs=[]
    for qt,n,m in PLANS[total]:
        if qt=='figure':
            figs=db.execute(f"select figure_image,count(*) c from ep_book_questions where chapter_id in ({ph}) and figure_image!='' and qtype in ('fill_blank','one_word','one_sentence','true_false','mcq') and length(text)<160 group by 1 having c>=2 order by c desc",ids).fetchall()
            figs=[f for f in figs if os.path.exists(FIG+f[0])]
            if not figs: continue
            fi=random.choice(figs[:6])[0]
            rows=db.execute(f"select bq_id,qtype,text,options_json,answer,pairs_json,figure_image from ep_book_questions where chapter_id in ({ph}) and figure_image=? and qtype in ('fill_blank','one_word','one_sentence','true_false') and length(text)<160",ids+[fi]).fetchall()
            random.shuffle(rows); rows=rows[:n]; used.update(r[0] for r in rows); secs.append(('figure',m,rows,fi)); continue
        rows=[]
        for t in [qt]+ALT.get(qt,[]):
            rows+=[r for r in db.execute(q,ids+[t]) if r[0] not in used]
            if len(rows)>=n*3: break
        random.shuffle(rows)
        if qt=='mcq': rows=[r for r in rows if r[3] and len(json.loads(r[3]))>=3]
        if qt=='match': rows=[r for r in rows if goodpairs(r[5])]
        if qt=='fill_blank': rows=[r for r in rows if r[4] and len(r[4])<25 and ('__' in r[2])]
        rows=rows[:n]
        if not rows: continue
        used.update(r[0] for r in rows); secs.append((qt,m,rows,None))
    tot=sum(s[1]*len(s[2]) for s in secs)
    while tot<total:
        need=total-tot; qt,m=('one_sentence',2) if need>=2 else ('one_word',1)
        rows=[r for r in db.execute(q,ids+[qt]) if r[0] not in used]
        if not rows: break
        r=random.choice(rows); used.add(r[0])
        for s in secs:
            if s[0]==qt and s[1]==m: s[2].append(r); break
        else: secs.append((qt,m,[r],None))
        tot+=m
    return chs,secs
def lang_of(sub,med): return 'hi' if sub=='Hindi' else ('en' if med=='English' else 'mr')
def D(n,lang): return str(n).translate(DEV) if lang!='en' else str(n)
def ins(qt,lang,rows):
    if qt=='figure': return {'mr':'चित्र पाहा व खालील प्रश्नांची उत्तरे लिहा.','hi':'चित्र देखकर नीचे दिए प्रश्नों के उत्तर लिखो।','en':'Look at the picture and answer the questions.'}[lang]
    if qt=='fill_blank':
        bank=[r[4] for r in rows]; random.shuffle(bank)
        return {'mr':'कंसातील योग्य शब्द निवडून रिकाम्या जागा भरा.','hi':'कोष्ठक में से उचित शब्द चुनकर रिक्त स्थानों की पूर्ति करो :','en':'Fill in the blanks with words from the bracket.'}[lang]+f' <span class="bank">( {" , ".join(html.escape(b) for b in bank)} )</span>'
    return INS[lang].get(qt,INS['mr'].get(qt,''))
def item_html(qt,r,lang,i):
    t=html.escape(r[2]); t=re.sub(r'_{3,}','<span class="dl"></span>',t)
    s=f'<div class="it"><span class="n">{D(i,lang)})</span> {t}'
    if qt=='mcq':
        ops=[o for o in json.loads(r[3]) if isinstance(o,str)][:4]
        s+='<div class="ops">'+''.join(f'<span>{(LETTERS[k] if lang!="en" else "abcd"[k])}) {html.escape(o)}</span>' for k,o in enumerate(ops))+'</div>'
    elif qt=='match':
        pairs=json.loads(r[5]); A=[str(p.get('a') or p.get('left')) for p in pairs]; B=[str(p.get('b') or p.get('right')) for p in pairs]; random.shuffle(B)
        s+='<table class="match"><tr><th>\'अ\' गट</th><th>उत्तर</th><th>\'ब\' गट</th></tr>'+''.join(f'<tr><td>{D(k+1,lang)}) {html.escape(a)}</td><td></td><td>{LETTERS[k]}) {html.escape(b)}</td></tr>' for k,(a,b) in enumerate(zip(A,B)))+'</table>'
    elif qt in ('true_false',): s+='<span class="br">( ............... )</span>'
    elif qt in ('one_word','vocabulary','odd_one','define'): s+='<div class="ln"></div>'
    elif qt in ('one_sentence',): s+='<div class="ln"></div><div class="ln"></div>'
    elif qt in ('short_answer','reason','explain','descriptive'): s+='<div class="ln"></div>'*3
    elif qt=='figure':
        if r[1] in ('one_sentence','one_word'): s+='<div class="ln"></div>'
        elif r[1]=='true_false': s+='<span class="br">( ............... )</span>'
    return s+'</div>'
def build(std,sub,med,total):
    db=sqlite3.connect(DB); chs,secs=pick(db,std,sub,med,total); lang=lang_of(sub,med)
    subj=sub if (med=='English' and sub!='Marathi') else MR[sub]
    title='संकलित मूल्यमापन (प्रथम सत्र) २०२६-२०२७' if lang!='en' else 'Summative Evaluation (First Term) 2026-2027'
    body='';qn=0;sub_i=0;prev_q=None
    for k,(qt,m,rows,fi) in enumerate(secs):
        if k%2==0: qn+=1; sub_i=0
        lab=(f'प्रश्न {D(qn,lang)}' if lang!='en' else f'Q.{qn}')+(f' {LETTERS[sub_i] if lang!="en" else "AB"[sub_i]})' if True else '')
        sub_i+=1
        body+=f'<div class="sec"><div class="sh"><span><b>{lab}</b> {ins(qt,lang,rows)}</span><span class="mk">({"गुण" if lang!="en" else "Marks"} {D(m*len(rows),lang)})</span></div>'
        if fi: body+=f'<div class="figwrap"><img src="data:image/jpeg;base64,{base64.b64encode(open(FIG+fi,"rb").read()).decode()}"></div>'
        body+=''.join(item_html(qt,r,lang,i+1) for i,r in enumerate(rows))+'</div>'
    tot=sum(s[1]*len(s[2]) for s in secs)
    H=f'''<!doctype html><html><head><meta charset="utf-8"><style>
@font-face{{font-family:M;src:url(file://{FONT})}} body{{font-family:M,sans-serif;font-size:12.5pt;margin:0;color:#000}}
.page{{padding:7mm 12mm 0}} .brand{{display:flex;align-items:center;gap:6px;font-size:9.5pt;color:#0b3d2e;margin-bottom:3px}} .brand img{{height:16px}} .brand span:last-child{{margin-left:auto}}
table.hdr{{width:100%;border:1.5px solid #000;border-collapse:collapse}} .hdr td{{padding:3px 8px;font-size:12.5pt}} .hdr .t{{text-align:center;font-size:15pt;font-weight:bold;border-bottom:1px solid #000}} .hdr .c{{text-align:center;font-weight:bold}}
.dot{{display:inline-block;border-bottom:1.5px dotted #000;min-width:120px;height:14px}} .dot.w{{width:62%}} .dot.m{{min-width:150px}} .box{{display:inline-block;border:1.5px solid #000;width:44px;height:20px;vertical-align:middle;text-align:center;font-weight:bold}}
.sec{{margin-top:9px;page-break-inside:avoid}} .sh{{display:flex;justify-content:space-between;gap:8px}} .mk{{white-space:nowrap}} .bank{{font-weight:normal}}
.it{{margin:4px 0 0 18px;position:relative}} .n{{display:inline-block;min-width:22px}} .dl{{display:inline-block;border-bottom:1.5px dotted #000;width:90px;height:12px}}
.ln{{border-bottom:1px dotted #555;height:19px;margin-left:22px}} .br{{float:right}} .ops{{display:flex;flex-wrap:wrap;gap:2px 20px;margin-left:22px;font-size:12pt}}
table.match{{border-collapse:collapse;margin:4px 0 0 22px;width:80%}} .match td,.match th{{border:1px solid #000;padding:2px 8px;font-size:12pt}} .match td:nth-child(2){{width:60px}}
.figwrap{{text-align:center;margin:4px 0}} .figwrap img{{max-height:170px;max-width:70%;border:1px solid #999}}
.sig{{display:flex;justify-content:space-between;margin-top:16px;border-top:1px dashed #777;padding-top:6px;font-size:11pt}}
</style></head><body><div class="page">
<div class="brand"><img src="file://{LOGO}"><b>TechGuruji स्मार्ट स्कूल हब</b><span>vijetaacademysangli.in/techguruji</span></div>
<table class="hdr"><tr><td colspan="3" class="t">{title}</td></tr>
<tr><td class="c">इयत्ता - {STD[std]}</td><td class="c" colspan="2">विषय - {subj} (लेखी)</td></tr>
<tr><td colspan="3">विद्यार्थ्याचे नाव - <span class="dot w"></span></td></tr>
<tr><td colspan="2">शाळेचे नाव - <span class="dot m"></span><span class="dot"></span></td><td>दिनांक - &nbsp; / &nbsp; /२०२६</td></tr>
<tr><td colspan="2">केंद्र - <span class="dot"></span> ता. <span class="dot"></span> जि. <span class="dot"></span></td><td>प्राप्त गुण : <span class="box"></span></td></tr>
<tr><td colspan="2">हजेरी क्रमांक - <span class="box"></span> &nbsp; तुकडी - ....... &nbsp; शिक्षकांची स्वाक्षरी - <span class="dot"></span></td><td>एकूण गुण <span class="box">{D(tot,lang)}</span></td></tr></table>
{body}
<div class="sig"><span>परीक्षक स्वाक्षरी: ______________</span><span>मुख्याध्यापक स्वाक्षरी व शिक्का</span></div>
</div></body></html>'''
    return H,tot,chs,secs,lang,subj,title
def make_docx(path,std,sub,med,secs,lang,subj,title,tot):
    d=Document(); st=d.styles['Normal']; st.font.name='Mukta'; st.font.size=Pt(11)
    for s in d.sections: s.left_margin=s.right_margin=Cm(1.5); s.top_margin=s.bottom_margin=Cm(1.2)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; r=p.add_run('TechGuruji स्मार्ट स्कूल हब | vijetaacademysangli.in/techguruji'); r.font.size=Pt(8)
    t=d.add_table(rows=6,cols=2); t.style='Table Grid'
    c=t.cell(0,0).merge(t.cell(0,1)); c.text=title; c.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER; c.paragraphs[0].runs[0].bold=True
    t.cell(1,0).text=f'इयत्ता - {STD[std]}'; t.cell(1,1).text=f'विषय - {subj} (लेखी)'
    t.cell(2,0).merge(t.cell(2,1)).text='विद्यार्थ्याचे नाव - ...............................................................'
    t.cell(3,0).text='शाळेचे नाव - ...............................................'; t.cell(3,1).text='दिनांक -    /    /२०२६'
    t.cell(4,0).text='केंद्र - ............... ता. ............... जि. ...............'; t.cell(4,1).text='प्राप्त गुण : [      ]'
    t.cell(5,0).text='हजेरी क्रमांक - [    ]  तुकडी - .......  शिक्षकांची स्वाक्षरी - ...............'; t.cell(5,1).text=f'एकूण गुण  {D(tot,lang)}'
    qn=0;sub_i=0
    for k,(qt,m,rows,fi) in enumerate(secs):
        if k%2==0: qn+=1; sub_i=0
        lab=(f'प्रश्न {D(qn,lang)}' if lang!='en' else f'Q.{qn}')+f' {LETTERS[sub_i] if lang!="en" else "AB"[sub_i]})'; sub_i+=1
        p=d.add_paragraph(); p.paragraph_format.space_before=Pt(8); r=p.add_run(f'{lab} '); r.bold=True
        p.add_run(re.sub('<[^>]+>','',ins(qt,lang,rows))+f'\t({"गुण" if lang!="en" else "Marks"} {D(m*len(rows),lang)})')
        if fi: d.add_picture(FIG+fi,height=Cm(5))
        for i,r in enumerate(rows):
            txt=re.sub(r'_{3,}','..............',r[2]); p=d.add_paragraph(f'{D(i+1,lang)}) {txt}'+('   ( ............ )' if (qt=='true_false' or (qt=='figure' and r[1]=='true_false')) else '')); p.paragraph_format.left_indent=Cm(0.6)
            if qt=='mcq': p=d.add_paragraph('    '.join(f'{(LETTERS[j] if lang!="en" else "abcd"[j])}) {o}' for j,o in enumerate([o for o in json.loads(r[3]) if isinstance(o,str)][:4]))); p.paragraph_format.left_indent=Cm(1.2)
            elif qt=='match':
                pairs=json.loads(r[5]); A=[str(x.get('a') or x.get('left')) for x in pairs]; B=[str(x.get('b') or x.get('right')) for x in pairs]; random.shuffle(B)
                mt=d.add_table(rows=len(A)+1,cols=3); mt.style='Table Grid'; mt.cell(0,0).text="'अ' गट"; mt.cell(0,1).text='उत्तर'; mt.cell(0,2).text="'ब' गट"
                for j,(a,b) in enumerate(zip(A,B)): mt.cell(j+1,0).text=f'{D(j+1,lang)}) {a}'; mt.cell(j+1,2).text=f'{LETTERS[j]}) {b}'
            n={'one_word':1,'vocabulary':1,'odd_one':1,'one_sentence':2,'short_answer':3,'reason':3,'explain':3,'descriptive':3}.get(qt if qt!='figure' else r[1],0)
            for _ in range(n): d.add_paragraph('.'*95).paragraph_format.left_indent=Cm(0.6)
    d.add_paragraph('\nपरीक्षक स्वाक्षरी: ______________\t\t\t\tमुख्याध्यापक स्वाक्षरी व शिक्का')
    d.save(path)
def sql_row(std,sub,med,secs,lang,subj,title,tot):
    sections=[];qn=0;sub_i=0
    for k,(qt,m,rows,fi) in enumerate(secs):
        if k%2==0: qn+=1; sub_i=0
        items=[]
        for r in rows:
            it={'bq_id':r[0],'text':r[2],'qtype':r[1],'source':'book','page_image':(r[6] or None),'show_image':bool(fi and r[6])}
            if r[4]: it['answer']=r[4]
            if r[1]=='mcq' and r[3]: it['options']=[o for o in json.loads(r[3]) if isinstance(o,str)]
            if r[1]=='match' and r[5]: it['pairs']=json.loads(r[5])
            items.append(it)
        sections.append({'q_no':qn,'sub':LETTERS[sub_i],'instruction':re.sub('<[^>]+>','',ins(qt,lang,rows)),'marks':m*len(rows),'qtype':qt if qt!='figure' else '','answer_space':'','oral':False,'chapter_ids':[],'items':items}); sub_i+=1
    meta={'exam_type':'sankalit','test_no':1,'standard':std,'subject':sub,'medium':med,'student_fields':True,'format':'lines','source':'book','oral_marks':0,'sections':sections}
    esc=lambda s:s.replace('\\','\\\\').replace("'","\\'")
    return f"INSERT INTO ep_papers (title, exam_name, std_label, exam_date, duration, total_marks, instructions, paper_json, paper_type) VALUES ('{esc(title+' - '+subj)}','प्रथम सत्र 2026-27','{STD[std]}',NULL,'',{tot},'','{esc(json.dumps(meta,ensure_ascii=False))}','assessment');\n"
async def render(jobs,sqlfile):
    sql=open(sqlfile,'w'); sql.write('-- TechGuruji प्रथम सत्र 2026-27 papers (assessment). Import into database `school`.\nSET NAMES utf8mb4;\n')
    async with async_playwright() as p:
        b=await p.chromium.launch(executable_path='/opt/.devin/google-chrome.sh'); pg=await b.new_page()
        for std,sub,med,total,out in jobs:
            H,tot,chs,secs,lang,subj,title=build(std,sub,med,total)
            await pg.set_content(H); await pg.wait_for_timeout(300)
            await pg.pdf(path=out,format='A4',print_background=True,margin={'top':'0','bottom':'13mm','left':'0','right':'0'},display_header_footer=True,header_template='<span></span>',
              footer_template='<div style="width:100%;margin:0 12mm;display:flex;justify-content:space-between;font-size:8px;color:#0b3d2e;border-top:1px solid #0b3d2e;padding-top:2px;font-family:sans-serif"><span>TechGuruji Smart School Hub - HPC | CCE | प्रश्नपत्रिका | टाचण | परिपाठ</span><span>vijetaacademysangli.in/techguruji &nbsp; <span class="pageNumber"></span>/<span class="totalPages"></span></span></div>')
            make_docx(out[:-4]+'.docx',std,sub,med,secs,lang,subj,title,tot); sql.write(sql_row(std,sub,med,secs,lang,subj,title,tot))
            print('ok',out,tot,[ (s[0],len(s[2])) for s in secs])
        await b.close()
    sql.close()
if __name__=='__main__':
    std,sub,med,total=int(sys.argv[1]),sys.argv[2],sys.argv[3],int(sys.argv[4])
    asyncio.run(render([(std,sub,med,total,sys.argv[5])],'/tmp/t.sql'))
