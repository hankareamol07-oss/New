"""Teacher-facing office files per chapter (and one combined workbook per book):
  * Excel quiz workbook – sheet 'MCQ': No | Topic | Question | A | B | C | D | Answer (A-D) | Answer text | Solution | Bloom | Page | Figure
                          sheet 'स्वाध्याय': No | Instruction | Type | Question | Options/Pairs | Answer | Marks | Bloom
                          sheet 'Answer key': compact No -> letter list;  sheet 'Notes' when notes exist
  * Word notes (.docx) – chapter heading, per topic: summary, key points, definitions, formulas/examples, names & dates,
    common mistakes, remember; textbook figures of the topic inline.
Pure python (openpyxl, python-docx); nothing is sent to an API."""
import json, os, re
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

LET = "ABCD"
HEAD = PatternFill("solid", fgColor="DDEBF7")
SAFE = re.compile(r"[\\/:*?\"<>|]+")


def _name(unit):
    return SAFE.sub("_", f"std{unit['std']}_{unit['subject']}_{unit['chapter_no']:02d}_{unit['chapter_title']}")[:80]


def _s(v, sep="\n"):
    """LLM fields may come back as str / list / dict – flatten to text for a cell."""
    if v is None:
        return ""
    if isinstance(v, dict):
        return sep.join(f"{k}: {_s(x, ' ')}" for k, x in v.items())
    if isinstance(v, list):
        return sep.join(_s(x, " ") for x in v)
    return str(v)


def _ans(v):
    if v is None:
        return ""
    try:
        j = json.loads(v)
        if isinstance(j, list):
            return ", ".join(map(str, j))
        if isinstance(j, dict):
            return "; ".join(f"{k}: {x}" for k, x in j.items())
    except (ValueError, TypeError):
        pass
    return str(v)


def _sheet(ws, header, rows, widths):
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), HEAD
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"


def _mcq_rows(store, unit, workdir):
    heads = {p["part_no"]: p["heading"] for p in store.parts(unit["unit_id"])}
    rows = []
    for i, q in enumerate(store.questions(unit["unit_id"], "mcq"), 1):
        opts = (json.loads(q["options_json"]) if q["options_json"] else []) + [""] * 4
        try:
            a = int(q["answer"])
        except (TypeError, ValueError):
            a = -1
        rows.append([i, heads.get(q["part_no"], ""), q["text"], *opts[:4], LET[a] if 0 <= a < 4 else "", opts[a] if 0 <= a < 4 else "",
                     q["explain"] or "", q["bloom"] or "", unit["first_page"], os.path.join(workdir, q["fig_file"]) if q["fig_file"] else ""])
    return rows


def _ex_rows(store, unit):
    rows = []
    for i, q in enumerate(store.questions(unit["unit_id"], "ex"), 1):
        op = json.loads(q["options_json"]) if q["options_json"] else None
        pr = json.loads(q["pairs_json"]) if q["pairs_json"] else None
        extra = " / ".join(op) if op else ("\n".join(f"{l} — {r}" for l, r in pr) if pr else "")
        rows.append([i, q["instruction"] or "", q["qtype"], q["text"], extra, _ans(q["answer"]), q["marks"] or 1, q["bloom"] or "", "generated" if q["source"] == "generated" else "textbook"])
    return rows


def _notes_rows(store, unit):
    heads = {p["part_no"]: p["heading"] for p in store.parts(unit["unit_id"])}
    rows = []
    for n in store.notes(unit["unit_id"]):
        d = json.loads(n["notes_json"])
        defs = d.get("definitions") or []
        rows.append([n["part_no"], heads.get(n["part_no"], ""), _s(d.get("summary")), "\n".join(f"• {_s(p, ' ')}" for p in (d.get("points") or [])),
                     "\n".join(f"{x.get('term')}: {_s(x.get('meaning'), ' ')}" if isinstance(x, dict) else _s(x, " ") for x in defs),
                     _s(d.get("formulas_examples")), _s(d.get("names_dates")), _s(d.get("mistakes")), _s(d.get("remember"))])
    return rows


def _fill_wb(wb, store, unit, workdir, first=True):
    mcq = _mcq_rows(store, unit, workdir)
    ws = wb.active if first else wb.create_sheet()
    ws.title = ("MCQ " + _name(unit))[:31] if not first else "MCQ"
    _sheet(ws, ["No", "Topic", "Question", "A", "B", "C", "D", "Answer", "Answer text", "Solution", "Bloom", "Page", "Figure file"], mcq,
           [5, 22, 60, 22, 22, 22, 22, 8, 22, 50, 12, 6, 30])
    _sheet(wb.create_sheet("स्वाध्याय"), ["No", "Instruction", "Type", "Question", "Options / Pairs", "Answer", "Marks", "Bloom", "Source"],
           _ex_rows(store, unit), [5, 30, 12, 60, 40, 50, 6, 12, 10])
    _sheet(wb.create_sheet("Answer key"), ["No", "Answer"], [[r[0], r[7]] for r in mcq], [6, 8])
    nr = _notes_rows(store, unit)
    if nr:
        _sheet(wb.create_sheet("Notes"), ["Part", "Topic", "Summary", "Key points", "Definitions", "Formulas / examples", "Names & dates", "Common mistakes", "Remember"],
               nr, [5, 22, 50, 60, 50, 40, 40, 40, 40])
    return len(mcq)


def export_excel(store, workdir, out_dir, units, per_book=True, log=print):
    os.makedirs(out_dir, exist_ok=True)
    books = {}
    for u in units:
        if not store.questions(u["unit_id"]):
            continue
        wb = Workbook()
        n = _fill_wb(wb, store, u, workdir)
        p = os.path.join(out_dir, _name(u) + ".xlsx")
        wb.save(p)
        books.setdefault(u["book_id"], []).append(u)
        log(f"  [excel] {os.path.basename(p)} ({n} MCQ)")
    if per_book:
        for bid, us in books.items():
            wb = Workbook()
            ws = wb.active
            ws.title = "MCQ"
            rows = []
            for u in us:
                rows += [[u["chapter_no"], u["chapter_title"]] + r for r in _mcq_rows(store, u, workdir)]
            _sheet(ws, ["Ch", "Chapter", "No", "Topic", "Question", "A", "B", "C", "D", "Answer", "Answer text", "Solution", "Bloom", "Page", "Figure file"], rows,
                   [4, 24, 5, 22, 60, 22, 22, 22, 22, 8, 22, 50, 12, 6, 30])
            exr = []
            for u in us:
                exr += [[u["chapter_no"], u["chapter_title"]] + r for r in _ex_rows(store, u)]
            _sheet(wb.create_sheet("स्वाध्याय"), ["Ch", "Chapter", "No", "Instruction", "Type", "Question", "Options / Pairs", "Answer", "Marks", "Bloom", "Source"], exr,
                   [4, 24, 5, 30, 12, 60, 40, 50, 6, 12, 10])
            _sheet(wb.create_sheet("Answer key"), ["Ch", "No", "Answer"], [[r[0], r[2], r[9]] for r in rows], [4, 6, 8])
            p = os.path.join(out_dir, SAFE.sub("_", f"std{us[0]['std']}_{us[0]['subject']}_ALL") + ".xlsx")
            wb.save(p)
            log(f"  [excel] {os.path.basename(p)} ({len(rows)} MCQ, {len(us)} chapters)")
    return len(books)


def export_docx(store, workdir, out_dir, units, log=print):
    from docx import Document
    from docx.shared import Inches, Pt
    os.makedirs(out_dir, exist_ok=True)
    n = 0
    for u in units:
        notes = store.notes(u["unit_id"])
        if not notes:
            continue
        heads = {p["part_no"]: p["heading"] for p in store.parts(u["unit_id"])}
        doc = Document()
        doc.styles["Normal"].font.size = Pt(12)
        doc.add_heading(f"इयत्ता {u['std']} – {u['subject']}", 0)
        doc.add_heading(f"{u['chapter_no']}. {u['chapter_title']}", 1)
        for nt in notes:
            d = json.loads(nt["notes_json"])
            doc.add_heading(heads.get(nt["part_no"], f"भाग {nt['part_no']}"), 2)
            if d.get("summary"):
                doc.add_paragraph(_s(d["summary"], " "))
            for title, key in (("महत्त्वाचे मुद्दे / Key points", "points"), ("सूत्रे व उदाहरणे / Formulas & examples", "formulas_examples"),
                               ("नावे, तारखा, मूल्ये / Names & dates", "names_dates"), ("सामान्य चुका / Common mistakes", "mistakes"), ("लक्षात ठेवा / Remember", "remember")):
                if d.get(key):
                    doc.add_paragraph(title).runs[0].bold = True
                    for x in (d[key] if isinstance(d[key], list) else [d[key]]):
                        doc.add_paragraph(_s(x, " "), style="List Bullet")
            if d.get("definitions"):
                doc.add_paragraph("व्याख्या / Definitions").runs[0].bold = True
                t = doc.add_table(rows=0, cols=2)
                t.style = "Table Grid"
                for x in (d["definitions"] if isinstance(d["definitions"], list) else [d["definitions"]]):
                    c = t.add_row().cells
                    c[0].text, c[1].text = (str(x.get("term", "")), _s(x.get("meaning"), " ")) if isinstance(x, dict) else (_s(x, " "), "")
            for f in store.figures(u["unit_id"], nt["part_no"])[:3]:
                fp = os.path.join(workdir, f["file"])
                if os.path.exists(fp):
                    try:
                        doc.add_picture(fp, width=Inches(4))
                        if f["caption"]:
                            doc.add_paragraph(f["caption"]).runs[0].italic = True
                    except Exception:  # noqa: BLE001
                        pass
        p = os.path.join(out_dir, _name(u) + "_notes.docx")
        doc.save(p)
        n += 1
        log(f"  [notes] {os.path.basename(p)}")
    return n


def export_paper_excel(store, workdir, out_dir, book, log=print):
    """One workbook per competitive paper: Question sheet (passage, question, images by position, A-D text + image, answer, solution) + Answer key."""
    items = store.paper_items(book["book_id"])
    if not items:
        return None
    os.makedirs(out_dir, exist_ok=True)
    name = SAFE.sub("_", f"paper_std{book['std'] or ''}_{os.path.splitext(os.path.basename(book['pdf']))[0]}")[:80]
    imgdir = os.path.join(out_dir, name + "_images")
    os.makedirs(imgdir, exist_ok=True)

    def cp(rel):
        if not rel:
            return ""
        src = os.path.join(workdir, rel)
        if not os.path.exists(src):
            return ""
        import shutil
        dst = os.path.join(imgdir, os.path.basename(rel))
        if not os.path.exists(dst):
            shutil.copy(src, dst)
        return os.path.basename(rel)

    rows = []
    for it in items:
        op = json.loads(it["options_json"] or "[]") + [""] * 4
        of = json.loads(it["option_files_json"] or "[]") + [None] * 4
        rows.append([it["page"], it["no"], it["section"] or "", it["topic"] or "", it["passage_id"] or "", it["passage"] or "", cp(it["passage_file"]), it["text"], cp(it["q_file"]),
                     op[0], cp(of[0]), op[1], cp(of[1]), op[2], cp(of[2]), op[3], cp(of[3]), it["answer"] or "", it["solution"] or "", it["marks"] or ""])
    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"
    _sheet(ws, ["Page", "Q.No", "Section", "Chapter / Topic", "Passage id", "Passage text", "Passage image", "Question", "Question image", "A", "A image", "B", "B image",
                "C", "C image", "D", "D image", "Answer", "Solution", "Marks"], rows,
           [6, 6, 16, 22, 9, 50, 22, 60, 24, 20, 22, 20, 22, 20, 22, 20, 22, 8, 40, 6])
    _sheet(wb.create_sheet("Answer key"), ["Q.No", "Section", "Answer"], [[r[1], r[2], r[17]] for r in rows], [6, 16, 8])
    chap = {}
    for r in rows:
        chap.setdefault((r[2], r[3]), []).append(str(r[1]))
    _sheet(wb.create_sheet("Chapters"), ["Section", "Chapter / Topic", "Questions", "Q.Nos"],
           [[k[0], k[1], len(v), ", ".join(v)] for k, v in chap.items()], [16, 30, 10, 60])
    _sheet(wb.create_sheet("Passages"), ["Page", "Passage id", "Text", "Image"], [[it["page"], it["passage_id"], it["passage"], cp(it["passage_file"])]
                                                                                for it in {(i["page"], i["passage_id"]): i for i in items if i["passage_id"]}.values()], [6, 9, 90, 22])
    p = os.path.join(out_dir, name + ".xlsx")
    wb.save(p)
    log(f"  [excel] {os.path.basename(p)} ({len(rows)} questions, images in {os.path.basename(imgdir)}\\)")
    return p
